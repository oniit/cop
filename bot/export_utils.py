import aiosqlite
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import datetime

async def generate_monthly_report(db_path, file_path):
    wb = Workbook()
    
    # Define styles
    header_fill = PatternFill(start_color="B4A7D6", end_color="B4A7D6", fill_type="solid") # Purple-ish
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    def apply_style(ws, start_row, start_col, end_row, end_col, is_header=False):
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = thin_border
                cell.alignment = center_align
                if is_header:
                    cell.fill = header_fill
                    cell.font = bold_font

    def auto_adjust_column_width(ws):
        for col_idx, col in enumerate(ws.columns, 1):
            max_length = 0
            column_letter = get_column_letter(col_idx)
            for cell in col:
                # Abaikan sel yang merupakan bagian dari merged cell (karena ukurannya menipu)
                is_merged = False
                for rng in ws.merged_cells.ranges:
                    if rng.min_row <= cell.row <= rng.max_row and rng.min_col <= cell.column <= rng.max_col:
                        is_merged = True
                        break
                if is_merged:
                    continue
                    
                try:
                    if cell.value:
                        lines = str(cell.value).split('\n')
                        for line in lines:
                            if len(line) > max_length:
                                max_length = len(line)
                except:
                    pass
                    
            adjusted_width = (max_length + 2)
            if adjusted_width < 6:
                adjusted_width = 6
            elif adjusted_width > 50:
                adjusted_width = 50
            ws.column_dimensions[column_letter].width = adjusted_width

    async with aiosqlite.connect(db_path) as db:
        # Sheet 1: Data Member
        ws1 = wb.active
        ws1.title = "DATA MEMBER"
        
        headers = ["NO", "NAMA LENGKAP", "CODE NAME", "AGENT ID", "MUSE", "TANGGAL LAHIR"]
        ws1.append(headers)
        
        async with db.execute("SELECT full_name, codename, agent_id, muse, dob FROM members") as cursor:
            members = await cursor.fetchall()
            for i, m in enumerate(members):
                ws1.append((i+1, m[0], m[1], m[2], m[3], m[4]))
                
        apply_style(ws1, 1, 1, 1, len(headers), True)

        # Retrieve members list for ordering
        async with db.execute("SELECT telegram_id, full_name, codename, agent_id FROM members") as cursor:
            all_members = await cursor.fetchall()
            
        current_month = datetime.datetime.now().strftime("%B %Y").upper()
            
        # Sheet 2: Attendance
        ws2 = wb.create_sheet("PRESENSI")
        ws2.merge_cells('B3:D3')
        ws2['B3'] = f"PRESENSI INTERNAL {current_month}"
        ws2['B3'].font = bold_font
        ws2['B3'].fill = header_fill
        
        # Headers
        ws2.merge_cells('B5:B6')
        ws2['B5'] = "NO"
        ws2.merge_cells('C5:C6')
        ws2['C5'] = "NAMA LENGKAP"
        ws2.merge_cells('D5:D6')
        ws2['D5'] = "CODE NAME"
        ws2.merge_cells('E5:E6')
        ws2['E5'] = "AGENT ID"

        # Fetch attendance sessions
        async with db.execute("SELECT session_id, title, created_at FROM attendance_sessions") as cursor:
            sessions = await cursor.fetchall()
            
        start_col = 6 # F
        if sessions:
            ws2.merge_cells(start_row=4, start_column=start_col, end_row=4, end_column=start_col+len(sessions)-1)
            ws2.cell(row=4, column=start_col).value = "PRESENSI KEGIATAN"
            for i, sess in enumerate(sessions):
                col = start_col + i
                ws2.cell(row=5, column=col).value = sess[1] # title
                date_str = sess[2].split(" ")[0] if sess[2] else ""
                ws2.cell(row=6, column=col).value = date_str
        
        end_col = start_col + len(sessions) - 1 if sessions else start_col
        ws2.merge_cells(start_row=5, start_column=end_col+1, end_row=6, end_column=end_col+1)
        ws2.cell(row=5, column=end_col+1).value = "TOTAL KEHADIRAN"
        ws2.merge_cells(start_row=5, start_column=end_col+2, end_row=6, end_column=end_col+2)
        ws2.cell(row=5, column=end_col+2).value = "KETERANGAN"
        
        apply_style(ws2, 4, 6, 4, end_col, True)
        apply_style(ws2, 5, 2, 6, end_col+2, True)
        
        # Data rows
        row_num = 7
        for i, mem in enumerate(all_members):
            tid, fname, cname, aid = mem
            ws2.cell(row=row_num, column=2, value=i+1)
            ws2.cell(row=row_num, column=3, value=fname)
            ws2.cell(row=row_num, column=4, value=cname)
            ws2.cell(row=row_num, column=5, value=aid)
            
            total = 0
            for j, sess in enumerate(sessions):
                col = start_col + j
                sid = sess[0]
                async with db.execute("SELECT 1 FROM attendance_logs WHERE session_id=? AND telegram_id=?", (sid, tid)) as c2:
                    is_present = await c2.fetchone()
                if is_present:
                    ws2.cell(row=row_num, column=col, value="V")
                    total += 1
                else:
                    ws2.cell(row=row_num, column=col, value="")
            ws2.cell(row=row_num, column=end_col+1, value=total)
            ws2.cell(row=row_num, column=end_col+2, value="")
            apply_style(ws2, row_num, 2, row_num, end_col+2, False)
            row_num += 1

        # Sheet 3: COP Points
        ws3 = wb.create_sheet("POINT COP")
        ws3.merge_cells('B2:E2')
        ws3['B2'] = f"POINT COP {current_month}"
        ws3['B2'].font = bold_font
        ws3['B2'].fill = header_fill
        
        ws3.merge_cells('B4:B6')
        ws3['B4'] = "NO"
        ws3.merge_cells('C4:C6')
        ws3['C4'] = "NAMA LENGKAP"
        ws3.merge_cells('D4:D6')
        ws3['D4'] = "CODE NAME"
        
        async with db.execute("SELECT DISTINCT reason, DATE(timestamp) as dt FROM point_logs WHERE point_type='COP'") as cursor:
            cop_activities = await cursor.fetchall()
            
        start_col = 5 # E
        if cop_activities:
            ws3.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=start_col+len(cop_activities)-1)
            ws3.cell(row=3, column=start_col).value = "KEGIATAN"
            for i, act in enumerate(cop_activities):
                col = start_col + i
                ws3.cell(row=4, column=col).value = "NAMA KEGIATAN"
                ws3.cell(row=5, column=col).value = act[0] # reason
                ws3.cell(row=6, column=col).value = act[1] # date
                
        end_col = start_col + len(cop_activities) - 1 if cop_activities else start_col
        ws3.merge_cells(start_row=4, start_column=end_col+1, end_row=6, end_column=end_col+1)
        ws3.cell(row=4, column=end_col+1).value = "TOTAL"
        ws3.merge_cells(start_row=4, start_column=end_col+2, end_row=6, end_column=end_col+2)
        ws3.cell(row=4, column=end_col+2).value = "KETERANGAN"
        
        if cop_activities:
            apply_style(ws3, 3, start_col, 3, end_col, True)
        apply_style(ws3, 4, 2, 6, end_col+2, True)
        
        row_num = 7
        for i, mem in enumerate(all_members):
            tid, fname, cname, aid = mem
            ws3.cell(row=row_num, column=2, value=i+1)
            ws3.cell(row=row_num, column=3, value=fname)
            ws3.cell(row=row_num, column=4, value=cname)
            
            total = 0
            for j, act in enumerate(cop_activities):
                col = start_col + j
                reason, dt = act
                async with db.execute("SELECT SUM(amount) FROM point_logs WHERE telegram_id=? AND point_type='COP' AND reason=?", (tid, reason)) as c2:
                    amt = await c2.fetchone()
                val = amt[0] if amt[0] else 0
                ws3.cell(row=row_num, column=col, value=val if val > 0 else "")
                total += val
            ws3.cell(row=row_num, column=end_col+1, value=total)
            ws3.cell(row=row_num, column=end_col+2, value="")
            apply_style(ws3, row_num, 2, row_num, end_col+2, False)
            row_num += 1

        # Sheet 4: Internal AIA Points
        ws_internal = wb.create_sheet("POINT INTERNAL")
        ws_internal['B2'] = f"BULAN {current_month}"
        ws_internal['B2'].font = bold_font
        
        row_offset = 4
        col_offset = 2 # B
        for i, mem in enumerate(all_members):
            tid, fname, cname, aid = mem
            
            ws_internal.merge_cells(start_row=row_offset, start_column=col_offset, end_row=row_offset, end_column=col_offset+2)
            ws_internal.cell(row=row_offset, column=col_offset).value = f"{fname} ({cname})"
            
            ws_internal.cell(row=row_offset+1, column=col_offset).value = "NO"
            ws_internal.cell(row=row_offset+1, column=col_offset+1).value = "JENIS KEGIATAN"
            ws_internal.cell(row=row_offset+1, column=col_offset+2).value = "POINT"
            
            apply_style(ws_internal, row_offset, col_offset, row_offset+1, col_offset+2, True)
            
            async with db.execute("SELECT event_name, points FROM aia_claims WHERE telegram_id=? AND claim_type='Internal' AND status='Approved'", (tid,)) as cursor:
                aia_pts = await cursor.fetchall()
                
            r = row_offset + 2
            total = 0
            for j, pt in enumerate(aia_pts):
                ws_internal.cell(row=r, column=col_offset).value = j+1
                ws_internal.cell(row=r, column=col_offset+1).value = pt[0]
                ws_internal.cell(row=r, column=col_offset+2).value = pt[1]
                total += pt[1]
                apply_style(ws_internal, r, col_offset, r, col_offset+2, False)
                r += 1
                
            if not aia_pts:
                ws_internal.cell(row=r, column=col_offset).value = 1
                ws_internal.cell(row=r, column=col_offset+1).value = ""
                ws_internal.cell(row=r, column=col_offset+2).value = ""
                apply_style(ws_internal, r, col_offset, r, col_offset+2, False)
                r += 1
                
            ws_internal.merge_cells(start_row=r, start_column=col_offset, end_row=r, end_column=col_offset+1)
            ws_internal.cell(row=r, column=col_offset).value = "TOTAL POINT"
            ws_internal.cell(row=r, column=col_offset+2).value = total
            apply_style(ws_internal, r, col_offset, r, col_offset+2, False)
            
            if i % 2 == 0:
                col_offset = 6 # F
            else:
                col_offset = 2
                row_offset = r + 2

        # Sheet 5: External AIA Points
        ws4 = wb.create_sheet("POINT EKSTERNAL")
        ws4['B2'] = f"BULAN {current_month}"
        ws4['B2'].font = bold_font
        
        row_offset = 4
        col_offset = 2 # B
        for i, mem in enumerate(all_members):
            tid, fname, cname, aid = mem
            
            ws4.merge_cells(start_row=row_offset, start_column=col_offset, end_row=row_offset, end_column=col_offset+2)
            ws4.cell(row=row_offset, column=col_offset).value = f"{fname} ({cname})"
            
            ws4.cell(row=row_offset+1, column=col_offset).value = "NO"
            ws4.cell(row=row_offset+1, column=col_offset+1).value = "JENIS KEGIATAN"
            ws4.cell(row=row_offset+1, column=col_offset+2).value = "POINT"
            
            apply_style(ws4, row_offset, col_offset, row_offset+1, col_offset+2, True)
            
            async with db.execute("SELECT event_name, points FROM aia_claims WHERE telegram_id=? AND claim_type='External' AND status='Approved'", (tid,)) as cursor:
                aia_pts = await cursor.fetchall()
                
            r = row_offset + 2
            total = 0
            for j, pt in enumerate(aia_pts):
                ws4.cell(row=r, column=col_offset).value = j+1
                ws4.cell(row=r, column=col_offset+1).value = pt[0]
                ws4.cell(row=r, column=col_offset+2).value = pt[1]
                total += pt[1]
                apply_style(ws4, r, col_offset, r, col_offset+2, False)
                r += 1
                
            if not aia_pts:
                ws4.cell(row=r, column=col_offset).value = 1
                ws4.cell(row=r, column=col_offset+1).value = ""
                ws4.cell(row=r, column=col_offset+2).value = ""
                apply_style(ws4, r, col_offset, r, col_offset+2, False)
                r += 1
                
            ws4.merge_cells(start_row=r, start_column=col_offset, end_row=r, end_column=col_offset+1)
            ws4.cell(row=r, column=col_offset).value = "TOTAL POINT"
            ws4.cell(row=r, column=col_offset+2).value = total
            apply_style(ws4, r, col_offset, r, col_offset+2, False)
            
            if i % 2 == 0:
                col_offset = 6 # F
            else:
                col_offset = 2
                row_offset = r + 2

        # Sheet 6: Data Perizinan
        ws5 = wb.create_sheet("PERIZINAN")
        ws5.merge_cells('F1:J1')
        ws5['F1'] = "DATA PERIZINAN CITY OF PRESTIGE"
        ws5['F1'].font = Font(bold=True, size=16)
        ws5['F1'].alignment = Alignment(horizontal="center")
        
        # We fetch all leave logs, joined with members
        async with db.execute("""
            SELECT l.leave_date, m.full_name, l.reason 
            FROM leave_logs l
            JOIN members m ON l.telegram_id = m.telegram_id
            ORDER BY l.id ASC
        """) as cursor:
            leaves = await cursor.fetchall()
            
        # We will just put them in a single mini table for the current month for simplicity,
        # or simulate the grid for demonstration.
        # Let's do a single table for "ALL TIME" since we didn't extract exact months from 'leave_date'
        
        ws5.merge_cells('B3:N3')
        ws5['B3'] = "PERIZINAN (KESELURUHAN)"
        ws5['B3'].fill = header_fill
        ws5['B3'].font = bold_font
        ws5['B3'].alignment = Alignment(horizontal="center")
        
        # Mini table 1
        start_row = 5
        start_col = 2 # B
        ws5.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col+3)
        ws5.cell(row=start_row, column=start_col).value = f"PERIZINAN {current_month}"
        
        ws5.cell(row=start_row+1, column=start_col).value = "NO"
        ws5.cell(row=start_row+1, column=start_col+1).value = "NAMA LENGKAP"
        ws5.cell(row=start_row+1, column=start_col+2).value = "TANGGAL PERIZINAN"
        ws5.cell(row=start_row+1, column=start_col+3).value = "KETERANGAN PERIZINAN"
        
        apply_style(ws5, start_row, start_col, start_row+1, start_col+3, True)
        
        r = start_row + 2
        for j, lv in enumerate(leaves):
            ws5.cell(row=r, column=start_col).value = j+1
            ws5.cell(row=r, column=start_col+1).value = lv[1] # full_name
            ws5.cell(row=r, column=start_col+2).value = lv[0] # date
            ws5.cell(row=r, column=start_col+3).value = lv[2] # reason
            apply_style(ws5, r, start_col, r, start_col+3, False)
            r += 1
            
        if not leaves:
            for c in range(4):
                ws5.cell(row=r, column=start_col+c).value = ""
            ws5.cell(row=r, column=start_col).value = 1
            apply_style(ws5, r, start_col, r, start_col+3, False)

        for sheet in wb.worksheets:
            auto_adjust_column_width(sheet)

    wb.save(file_path)
